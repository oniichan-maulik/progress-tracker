"""Turn tasks into tables, summaries and styled Excel workbooks."""

from dataclasses import dataclass
from datetime import date
from io import BytesIO

import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from core.models import COMPLETED, IN_PROGRESS, NOT_STARTED, UNDER_REVIEW, Task
from core.services import is_overdue, month_label, week_of_month, week_range

COLUMNS = ["Date", "Week", "Task", "Status", "Hours", "Notes", "Assigned By", "Due Date"]

STATUS_FILLS = {
    NOT_STARTED: "E7E6E2",
    IN_PROGRESS: "CDE2FB",
    UNDER_REVIEW: "FDE7B0",
    COMPLETED: "C9EDC9",
}


def tasks_to_df(tasks: list[Task]) -> pd.DataFrame:
    rows = [
        {
            "id": t.id,
            "Date": t.work_date,
            "Week": week_of_month(t.work_date),
            "Task": t.title,
            "Status": t.status,
            "Hours": t.hours,
            "Notes": t.notes,
            "Assigned By": t.assigner.full_name if t.assigner else "",
            "Due Date": t.due_date,
            "Overdue": is_overdue(t),
            "Deletion Requested": t.deletion_requested_at is not None,
            "Employee": t.owner.full_name,
            "user_id": t.user_id,
        }
        for t in tasks
    ]
    df = pd.DataFrame(rows, columns=["id", *COLUMNS, "Overdue", "Deletion Requested", "Employee", "user_id"])
    df["Hours"] = pd.to_numeric(df["Hours"], errors="coerce")
    return df


def summarize(df: pd.DataFrame) -> dict:
    total = len(df)
    completed = int((df["Status"] == COMPLETED).sum()) if total else 0
    return {
        "tasks": total,
        "completed": completed,
        "in_progress": int((df["Status"] == IN_PROGRESS).sum()) if total else 0,
        "under_review": int((df["Status"] == UNDER_REVIEW).sum()) if total else 0,
        "not_started": int((df["Status"] == NOT_STARTED).sum()) if total else 0,
        "hours": float(df["Hours"].fillna(0).sum()) if total else 0.0,
        "completion_rate": completed / total if total else 0.0,
    }


def weekly_breakdown(df: pd.DataFrame, month: str) -> list[tuple[int, date, date, pd.DataFrame]]:
    """[(week number, start, end, that week's rows)] for weeks that have tasks."""
    out = []
    for week, rows in df.sort_values(["Date", "id"]).groupby("Week", sort=True):
        start, end = week_range(month, int(week))
        out.append((int(week), start, end, rows))
    return out


def fmt_hours(h: float) -> str:
    return f"{h:g} h"


# ---------------------------------------------------------------- Excel


@dataclass
class ReportSection:
    employee: str
    month: str
    review_state: str
    df: pd.DataFrame
    reviewer_comment: str = ""


EXPORT_COLUMNS = ["Date", "Task", "Status", "Hours", "Notes", "Assigned By", "Due Date"]
COLUMN_WIDTHS = [12, 48, 15, 9, 40, 18, 12]

_thin = Side(style="thin", color="C8C7C2")
_BORDER = Border(left=_thin, right=_thin, top=_thin, bottom=_thin)
_HEADER_FILL = PatternFill("solid", fgColor="1C5CAB")
_WEEK_FILL = PatternFill("solid", fgColor="F0EFEC")
_SUBTOTAL_FILL = PatternFill("solid", fgColor="FAFAF8")


def _sheet_title(name: str, used: set[str]) -> str:
    base = "".join(c for c in name if c not in r"[]:*?/\\")[:31] or "Report"
    title, n = base, 2
    while title in used:
        suffix = f" ({n})"
        title, n = base[: 31 - len(suffix)] + suffix, n + 1
    used.add(title)
    return title


def _write_section(ws, section: ReportSection, org: str) -> None:
    ncols = len(EXPORT_COLUMNS)
    stats = summarize(section.df)

    ws["A1"] = f"{org} - Monthly Progress Report"
    ws["A1"].font = Font(size=14, bold=True)
    ws["A2"] = f"Employee: {section.employee}"
    ws["A3"] = f"Month: {month_label(section.month)}"
    ws["A4"] = f"Review status: {section.review_state}"
    if section.reviewer_comment:
        ws["A5"] = f"Reviewer comment: {section.reviewer_comment}"

    summary = [
        ("Total tasks", stats["tasks"]),
        ("Completed", stats["completed"]),
        ("Total hours", stats["hours"]),
        ("Completion rate", stats["completion_rate"]),
    ]
    for i, (label, value) in enumerate(summary):
        ws.cell(row=2 + i, column=5, value=label).font = Font(bold=True)
        cell = ws.cell(row=2 + i, column=6, value=value)
        if label == "Completion rate":
            cell.number_format = "0%"

    row = 7
    for col, name in enumerate(EXPORT_COLUMNS, start=1):
        cell = ws.cell(row=row, column=col, value=name)
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = _HEADER_FILL
        cell.border = _BORDER
        cell.alignment = Alignment(horizontal="center", vertical="center")
    ws.freeze_panes = ws.cell(row=row + 1, column=1)
    row += 1

    if section.df.empty:
        ws.cell(row=row, column=1, value="No tasks logged this month.").font = Font(italic=True)

    for week, start, end, rows in weekly_breakdown(section.df, section.month):
        ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=ncols)
        cell = ws.cell(row=row, column=1, value=f"Week {week}  ({start:%d %b} - {end:%d %b})")
        cell.font = Font(bold=True)
        cell.fill = _WEEK_FILL
        row += 1

        for _, r in rows.iterrows():
            values = [
                r["Date"],
                r["Task"],
                r["Status"],
                None if pd.isna(r["Hours"]) else float(r["Hours"]),
                r["Notes"],
                r["Assigned By"],
                r["Due Date"],
            ]
            for col, value in enumerate(values, start=1):
                cell = ws.cell(row=row, column=col, value=value)
                cell.border = _BORDER
                cell.alignment = Alignment(vertical="top", wrap_text=col in (2, 5))
                if isinstance(value, date):
                    cell.number_format = "DD-MMM-YYYY"
            status_cell = ws.cell(row=row, column=3)
            status_cell.fill = PatternFill("solid", fgColor=STATUS_FILLS.get(r["Status"], "FFFFFF"))
            status_cell.alignment = Alignment(horizontal="center", vertical="top")
            row += 1

        week_stats = summarize(rows)
        ws.cell(row=row, column=2, value=f"Week {week} subtotal - {week_stats['completed']}/{week_stats['tasks']} completed")
        ws.cell(row=row, column=4, value=week_stats["hours"])
        for col in range(1, ncols + 1):
            c = ws.cell(row=row, column=col)
            c.font = Font(bold=True)
            c.fill = _SUBTOTAL_FILL
            c.border = _BORDER
        row += 2

    if not section.df.empty:
        ws.cell(row=row, column=2, value="Month total").font = Font(bold=True, size=12)
        ws.cell(row=row, column=4, value=stats["hours"]).font = Font(bold=True, size=12)

    for i, width in enumerate(COLUMN_WIDTHS, start=1):
        ws.column_dimensions[get_column_letter(i)].width = width


def build_workbook(sections: list[ReportSection], org: str) -> bytes:
    """One sheet per section (an employee's month). Returns .xlsx bytes."""
    wb = Workbook()
    wb.remove(wb.active)
    used: set[str] = set()
    for section in sections:
        name = section.employee if len(sections) > 1 else month_label(section.month)
        ws = wb.create_sheet(_sheet_title(name, used))
        _write_section(ws, section, org)
    if not sections:
        wb.create_sheet("Report")["A1"] = "No data."
    buf = BytesIO()
    wb.save(buf)
    return buf.getvalue()
