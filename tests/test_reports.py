from datetime import date
from io import BytesIO

from openpyxl import load_workbook

from core.models import COMPLETED, IN_PROGRESS
from core.reports import ReportSection, build_workbook, summarize, tasks_to_df, weekly_breakdown
from core.services import add_task, assign_task, get_tasks

OCT = "2026-10"


def _df(session, admin, alice):
    add_task(session, alice.id, date(2026, 10, 2), "Setup", COMPLETED, hours=2.5)
    add_task(session, alice.id, date(2026, 10, 6), "Build", IN_PROGRESS, hours=1)
    assign_task(session, admin, [alice.id], "Audit", due_date=date(2026, 10, 30), today=date(2026, 10, 7))
    return tasks_to_df(get_tasks(session, user_id=alice.id, month=OCT))


def test_summary_and_weeks(session, admin, alice):
    df = _df(session, admin, alice)
    stats = summarize(df)
    assert stats["tasks"] == 3 and stats["completed"] == 1 and stats["hours"] == 3.5
    assert round(stats["completion_rate"], 2) == 0.33
    weeks = weekly_breakdown(df, OCT)
    assert [(w, len(rows)) for w, _, _, rows in weeks] == [(1, 1), (2, 2)]
    assert df.loc[df["Task"] == "Audit", "Assigned By"].item() == "Bea Boss"


def test_empty_summary():
    stats = summarize(tasks_to_df([]))
    assert stats["tasks"] == 0 and stats["hours"] == 0 and stats["completion_rate"] == 0


def test_excel_export(session, admin, alice, bob):
    df = _df(session, admin, alice)
    single = load_workbook(BytesIO(build_workbook([ReportSection("Alice Archer", OCT, "Draft", df)], "Acme")))
    assert single.sheetnames == ["October 2026"]
    ws = single.active
    values = [c for row in ws.iter_rows(values_only=True) for c in row if c is not None]
    assert "Acme - Monthly Progress Report" in values
    assert "Setup" in values and "Audit" in values
    assert any(isinstance(v, str) and v.startswith("Week 2") for v in values)
    assert 3.5 in values

    empty = tasks_to_df([])
    multi = load_workbook(
        BytesIO(
            build_workbook(
                [ReportSection("Alice Archer", OCT, "Draft", df), ReportSection("Bob Baker", OCT, "Draft", empty)], "Acme"
            )
        )
    )
    assert multi.sheetnames == ["Alice Archer", "Bob Baker"]
